"""Drawing primitives for the leadership briefing, on top of python-pptx.

The deck was first generated with a JavaScript library in a scratch directory,
and that directory was lost with the session that held it. This module is the
same design rebuilt in Python and kept in the repository. The repository allows
no npm step, and the organisation must be able to rebuild the deck with only
the development requirements installed. Every measurement here (inches), colour
and type size is carried over from that generator, so the output matches the
deck it replaces slide for slide.

Conventions the slide code relies on:

* Positions and sizes are **inches** on a 13.333 × 7.5 in (16:9) slide.
* Text is written with a small markup: ``**bold**``, ``{{accent}}`` (bold, in
  the accent colour) and ``[[highlight]]`` (for an interim figure; unused in
  the final deck). ``<<name>>`` is replaced from ``Deck.values`` first, so a
  number that comes from ``facts.json`` is never typed into slide text by hand.
* Every shape carries a 1 pt outline in its own fill colour unless a line is
  given. The original generator did this, and the slides were laid out with it.
"""

from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- palette & type
NAVY, NAVY2, SLATE, INK, MUTED = "14213D", "2B3F66", "5B6B82", "2A2F36", "6B7380"
LINE, CARD, ICE, WHITE, ACCENT = "D3D9E1", "F2F4F7", "C9D3E3", "FFFFFF", "C8553D"
ACCENT_T, HL, TINT = "F7E3DE", "FFE58A", "E8EDF5"
CLASS_COLOURS = {"a": "14213D", "b": "3D6A9E", "c": "8D99AE", "d": ACCENT}  # coverage classes
FONT = "Calibri"
W, H, M = 13.333, 7.5, 0.5
CW = W - 2 * M

SECTIONS = [("A", "Why this exists"), ("B", "What we built"), ("C", "How it works"),
            ("D", "Evidence"), ("E", "What is not done"), ("F", "Road to production")]

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_TOKEN = re.compile(r"(\*\*[^*]+\*\*|\[\[[^\]]+\]\]|\{\{[^}]+\}\})")
_VALUE = re.compile(r"<<([a-zA-Z0-9_]+)>>")


def rgb(hex6: str) -> RGBColor:
	return RGBColor.from_string(hex6)


class Deck:
	"""One presentation being built: the file, the slide counter and the values
	that ``<<name>>`` placeholders in slide text resolve to."""

	def __init__(self, repo: Path, facts: dict, values: dict, shots: dict):
		self.repo = repo
		self.facts = facts
		self.values = values
		self.shots = shots
		self.prs = Presentation()
		self.prs.slide_width = Inches(W)
		self.prs.slide_height = Inches(H)
		self.blank = self.prs.slide_layouts[6]
		self.slide_no = 0
		self.missing_shots: list[str] = []

	# ------------------------------------------------------------ text
	def fill(self, text: str) -> str:
		"""Resolve ``<<name>>`` from the computed values. An unknown name is a
		build error, not a silent gap on a slide."""
		def sub(m):
			key = m.group(1)
			if key not in self.values:
				raise KeyError(f"slide text names <<{key}>>, which facts.json does not provide")
			return str(self.values[key])
		return _VALUE.sub(sub, text)

	def runs(self, text: str, base: dict) -> list[tuple[str, dict]]:
		text = self.fill(text)
		out, last = [], 0
		for m in _TOKEN.finditer(text):
			if m.start() > last:
				out.append((text[last:m.start()], dict(base)))
			tok = m.group(0)
			if tok.startswith("**"):
				out.append((tok[2:-2], {**base, "bold": True}))
			elif tok.startswith("[["):
				out.append((tok[2:-2], {**base, "highlight": HL, "bold": True, "color": INK}))
			else:
				out.append((tok[2:-2], {**base, "color": ACCENT, "bold": True}))
			last = m.end()
		if last < len(text):
			out.append((text[last:], dict(base)))
		return out or [("", dict(base))]

	def new_slide(self, background: str = WHITE):
		s = self.prs.slides.add_slide(self.blank)
		self.slide_no += 1
		bg = s.background.fill
		bg.solid()
		bg.fore_color.rgb = rgb(background)
		return s


# ---------------------------------------------------------------- low-level XML helpers
def _set_run(run, o: dict) -> None:
	f = run.font
	f.name = o.get("font", FONT)
	f.size = Pt(o.get("size", 12))
	f.bold = bool(o.get("bold"))
	if o.get("italic"):
		f.italic = True
	f.color.rgb = rgb(o.get("color", INK))
	rpr = run._r.get_or_add_rPr()
	if o.get("spacing"):
		rpr.set("spc", str(int(o["spacing"] * 100)))
	if o.get("highlight"):
		hl = etree.SubElement(rpr, qn("a:highlight"))
		etree.SubElement(hl, qn("a:srgbClr")).set("val", o["highlight"])


def _para_format(p, o: dict) -> None:
	p.alignment = _ALIGN[o.get("align", "left")]
	ppr = p._p.get_or_add_pPr()
	if o.get("bullet"):
		ppr.set("marL", "177800")
		ppr.set("indent", "-177800")
	else:
		ppr.set("marL", "0")
		ppr.set("indent", "0")
	if o.get("lsm"):
		p.line_spacing = o["lsm"]
	if o.get("gap") is not None:
		p.space_after = Pt(o["gap"])
	for tag in ("a:buNone", "a:buChar", "a:buSzPct"):
		for el in ppr.findall(qn(tag)):
			ppr.remove(el)
	if o.get("bullet"):
		etree.SubElement(ppr, qn("a:buSzPct")).set("val", "100000")
		etree.SubElement(ppr, qn("a:buChar")).set("char", "•")
	else:
		etree.SubElement(ppr, qn("a:buNone"))


def _line(shape, colour: str | None, width: float | None, dash: str | None = None) -> None:
	if colour is None:
		shape.line.fill.background()
		return
	shape.line.color.rgb = rgb(colour)
	shape.line.width = Pt(width or 1)
	if dash:
		ln = shape._element.spPr.find(qn("a:ln"))
		etree.SubElement(ln, qn("a:prstDash")).set("val", dash)


# ---------------------------------------------------------------- text boxes
def text(d: Deck, s, paragraphs: list[list[tuple[str, dict]]], x, y, w, h, *, valign="top",
         margin=0.0, fill: str | None = None, line: str | None = None, align="left", lsm=None,
         para_opts: list[dict] | None = None):
	"""A text box holding ``paragraphs``: each a list of (text, run options)."""
	tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
	tf = tb.text_frame
	tf.word_wrap = True
	tf.auto_size = MSO_AUTO_SIZE.NONE
	inset = Pt(margin) if margin else Emu(0)
	tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = inset
	tf.vertical_anchor = _ANCHOR[valign]
	if fill:
		tb.fill.solid()
		tb.fill.fore_color.rgb = rgb(fill)
	if line:
		tb.line.color.rgb = rgb(line)
		tb.line.width = Pt(0.75)
	for i, runs in enumerate(paragraphs):
		p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
		po = dict(para_opts[i]) if para_opts else {}
		po.setdefault("align", align)
		if lsm:
			po["lsm"] = lsm
		_para_format(p, po)
		for t, o in runs:
			r = p.add_run()
			r.text = t
			_set_run(r, o)
	return tb


def tx(d: Deck, s, content, x, y, w, h, *, base: dict | None = None, valign="top", align="left",
       margin=0.0, fill=None, line=None, lsm=None):
	"""The workhorse. ``content`` is a string (one paragraph of marked-up runs)
	or a list of paragraphs: strings or dicts {t, bullet, size, color, bold,
	italic, gap}. A list paragraph without ``gap`` has 4 pt after it."""
	b = {"font": FONT, "size": 12, "color": INK}
	b.update(base or {})
	if isinstance(content, str):
		return text(d, s, [d.runs(content, b)], x, y, w, h, valign=valign, margin=margin, fill=fill,
		            line=line, align=align, lsm=lsm)
	paras, popts = [], []
	for item in content:
		o = {"t": item} if isinstance(item, str) else item
		ro = dict(b)
		for k_src, k_dst in (("size", "size"), ("color", "color")):
			if o.get(k_src):
				ro[k_dst] = o[k_src]
		if o.get("bold"):
			ro["bold"] = True
		if o.get("italic"):
			ro["italic"] = True
		paras.append(d.runs(o["t"], ro))
		popts.append({"bullet": o.get("bullet"), "gap": o.get("gap", b.get("gap", 4))})
	return text(d, s, paras, x, y, w, h, valign=valign, margin=margin, fill=fill, line=line,
	            align=align, lsm=lsm, para_opts=popts)


def label(d: Deck, s, runs: list[tuple[str, dict]], x, y, w, h, *, valign="middle", align="left",
          margin=0.0, paragraphs: list[list[tuple[str, dict]]] | None = None):
	"""A text box written from explicit runs (no markup), for the few places the
	original placed text directly rather than through ``tx``."""
	paras = paragraphs if paragraphs is not None else [runs]
	full = []
	for pr in paras:
		full.append([(t, {"font": FONT, **o}) for t, o in pr])
	return text(d, s, full, x, y, w, h, valign=valign, margin=margin, align=align)


# ---------------------------------------------------------------- shapes
_SHAPES = {
	"rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL, "chevron": MSO_SHAPE.CHEVRON,
	"pentagon": MSO_SHAPE.PENTAGON, "right_arrow": MSO_SHAPE.RIGHT_ARROW,
}


def shape(s, kind: str, x, y, w, h, fill: str | None, line: str | None = "same", line_w: float | None = None,
          dash: str | None = None):
	sp = s.shapes.add_shape(_SHAPES[kind], Inches(x), Inches(y), Inches(w), Inches(h))
	if fill:
		sp.fill.solid()
		sp.fill.fore_color.rgb = rgb(fill)
	else:
		sp.fill.background()
	_line(sp, fill if line == "same" else line, line_w, dash)
	_unstyle(sp)
	return sp


def _unstyle(sp) -> None:
	"""Drop the theme style python-pptx attaches to a new shape or connector.
	Its effect reference draws a drop shadow the original shapes did not have;
	fill and line are always set explicitly here, so nothing else is lost."""
	st = sp._element.find(qn("p:style"))
	if st is not None:
		sp._element.remove(st)


def rule(s, x, y, w, h, colour: str, width: float, dash: str | None = None):
	"""A straight line from (x, y) to (x + w, y + h)."""
	c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y + h))
	c.line.color.rgb = rgb(colour)
	c.line.width = Pt(width)
	_unstyle(c)
	if dash:
		ln = c._element.spPr.find(qn("a:ln"))
		etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
	return c


def card(s, x, y, w, h, fill: str = CARD):
	return shape(s, "rect", x, y, w, h, fill)


def badge(d: Deck, s, n, x, y, dia=0.34, fill=ACCENT):
	shape(s, "oval", x, y, dia, dia, fill)
	label(d, s, [(str(n), {"size": 12, "bold": True, "color": WHITE})], x, y, dia, dia, align="center")


def kpi(d: Deck, s, x, y, w, value: str, caption: str, size=36, colour=NAVY, lh=0.6):
	text(d, s, [d.runs(value, {"font": FONT, "size": size, "bold": True, "color": colour})], x, y, w, 0.7, valign="bottom")
	text(d, s, [d.runs(caption, {"font": FONT, "size": 11, "color": MUTED})], x, y + 0.74, w, lh, valign="top")


# ---------------------------------------------------------------- frames
def footer(d: Deck, s, source: str | None, dark: bool = False):
	if source:
		text(d, s, [d.runs(source, {"font": FONT, "size": 9, "color": ICE if dark else MUTED})], M, 6.98, 11.2, 0.36)
	label(d, s, [("Consilium  |  ", {"size": 9, "color": ICE if dark else MUTED}),
	             (str(d.slide_no), {"size": 9, "bold": True, "color": WHITE if dark else NAVY})],
	      W - M - 1.6, 6.98, 1.6, 0.3, align="right")


def frame(d: Deck, title: str, section: str | None, source: str | None, notes: str | None):
	s = d.new_slide()
	if section:
		sec = next((t for k, t in SECTIONS if k == section), None)
		lab = f"{section}  ·  {sec.upper()}" if sec else section.upper()
		label(d, s, [(lab, {"size": 10, "color": MUTED, "spacing": 1})], M, 0.28, 8, 0.28)
	text(d, s, [d.runs(title, {"font": FONT, "size": 22, "bold": True, "color": NAVY})], M, 0.6, CW, 0.95,
	     valign="top", lsm=0.95)
	footer(d, s, source)
	notes_(d, s, notes)
	return s


def notes_(d: Deck, s, notes: str | None):
	if notes:
		s.notes_slide.notes_text_frame.text = d.fill(notes)


def divider(d: Deck, letter: str, title: str, question: str, notes: str):
	s = d.new_slide(NAVY)
	label(d, s, [(letter, {"size": 150, "bold": True, "color": ACCENT})], M, 1.55, 2.2, 2.2, valign="top")
	label(d, s, [(title, {"size": 36, "bold": True, "color": WHITE})], 3.0, 2.0, 9.5, 0.9)
	label(d, s, [(question, {"size": 18, "color": ICE})], 3.0, 2.95, 9.3, 1.2, valign="top")
	tw = CW / len(SECTIONS)
	for i, (k, t) in enumerate(SECTIONS):
		cur = k == letter
		label(d, s, [(k + "  ", {"size": 12, "bold": True, "color": ACCENT if cur else "8391A8"}),
		             (t, {"size": 12, "bold": cur, "color": WHITE if cur else "8391A8"})],
		      M + i * tw, 5.9, tw - 0.1, 0.35)
		shape(s, "rect", M + i * tw, 5.8, tw - 0.15, 0.04, ACCENT if cur else "3A4A6A")
	footer(d, s, None, dark=True)
	notes_(d, s, notes)
	return s


# ---------------------------------------------------------------- tables
def table(d: Deck, s, rows: list[list], x, y, w, col_w: list[float], *, font_size=11, h_size=None,
          row_h=None, zebra=False, header=True):
	"""A table with a navy header row, a thin rule under each row and nothing
	else. A cell is a string (markup allowed) or a dict {t, bold, align, color,
	fill, size}."""
	n_rows, n_cols = len(rows), len(col_w)
	heights = row_h if isinstance(row_h, list) else [row_h or 0.4] * n_rows
	gf = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(sum(heights)))
	tbl = gf.table
	# No theme table style: the original tables were plain.
	tbl_pr = tbl._tbl.tblPr
	for attr in ("firstRow", "bandRow"):
		if attr in tbl_pr.attrib:
			del tbl_pr.attrib[attr]
	sid = tbl_pr.find(qn("a:tableStyleId"))
	if sid is not None:
		tbl_pr.remove(sid)
	for j, cw in enumerate(col_w):
		tbl.columns[j].width = Inches(cw)
	for i, hgt in enumerate(heights):
		tbl.rows[i].height = Inches(hgt)
	for i, row in enumerate(rows):
		for j, raw in enumerate(row):
			c = {"t": raw} if isinstance(raw, str) else dict(raw)
			hdr = i == 0 and header
			cell = tbl.cell(i, j)
			size = (h_size or font_size) if hdr else c.get("size", font_size)
			fill = NAVY if hdr else (c.get("fill") or ("F7F8FA" if zebra and i % 2 == 0 else WHITE))
			base = {"font": FONT, "size": size, "color": WHITE if hdr else c.get("color", INK),
			        "bold": bool(hdr or c.get("bold"))}
			runs = [(d.fill(c["t"]), base)] if hdr else d.runs(c["t"], base)
			tf = cell.text_frame
			tf.word_wrap = True
			p = tf.paragraphs[0]
			_para_format(p, {"align": c.get("align", "left")})
			for t, o in runs:
				r = p.add_run()
				r.text = t
				_set_run(r, o)
			cell.margin_left = cell.margin_right = Pt(5)
			cell.margin_top = cell.margin_bottom = Pt(3)
			cell.vertical_anchor = MSO_ANCHOR.MIDDLE
			tc_pr = cell._tc.get_or_add_tcPr()
			for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
				ln = etree.SubElement(tc_pr, qn(side))
				if side == "a:lnB":
					ln.set("w", "9525")
					sf = etree.SubElement(ln, qn("a:solidFill"))
					etree.SubElement(sf, qn("a:srgbClr")).set("val", LINE)
				else:
					ln.set("w", "0")
					etree.SubElement(ln, qn("a:noFill"))
			sf = etree.SubElement(tc_pr, qn("a:solidFill"))
			etree.SubElement(sf, qn("a:srgbClr")).set("val", fill)
	return gf


# ---------------------------------------------------------------- pictures
def shot(d: Deck, s, key: str, x, y, w, h, caption: str | None = None):
	"""A screenshot from docs/guides/images, scaled to fit the box (top-aligned,
	centred across), in a white frame with a soft shadow. A missing file draws a
	labelled placeholder and is reported by the build."""
	spec = d.shots.get(key)
	path = d.repo / spec["file"] if spec else None
	if path and path.exists():
		with Image.open(path) as im:
			iw, ih = im.size
		crop = spec.get("crop")  # [x, y, w, h] in pixels
		cw_, ch_ = (crop[2], crop[3]) if crop else (iw, ih)
		cw_, ch_ = min(cw_, iw - (crop[0] if crop else 0)), min(ch_, ih - (crop[1] if crop else 0))
		r = min(w / cw_, h / ch_)
		dw, dh = cw_ * r, ch_ * r
		dx, dy = x + (w - dw) / 2, y
		fr = shape(s, "rect", dx - 0.02, dy - 0.02, dw + 0.04, dh + 0.04, WHITE, line=LINE, line_w=0.75)
		_shadow(fr)
		pic = s.shapes.add_picture(str(path), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
		if crop:
			pic.crop_left = crop[0] / iw
			pic.crop_top = crop[1] / ih
			pic.crop_right = max(0.0, 1 - (crop[0] + cw_) / iw)
			pic.crop_bottom = max(0.0, 1 - (crop[1] + ch_) / ih)
		pic._element.nvPicPr.cNvPr.set("descr", caption or spec.get("desc", key))
		if caption:
			tx(d, s, caption, x, y + h + 0.06, w, 0.25, base={"size": 9, "color": MUTED, "italic": True})
		return
	d.missing_shots.append(key)
	shape(s, "rect", x, y, w, h, "F7F8FA", line="B7C0CC", line_w=1, dash="dash")
	shape(s, "rect", x, y, w, 0.3, "E4E8EE")
	for i in range(3):
		shape(s, "oval", x + 0.14 + i * 0.18, y + 0.09, 0.12, 0.12, "B7C0CC")
	tx(d, s, [{"t": "SCREENSHOT MISSING", "size": 11, "bold": True, "color": SLATE, "gap": 6},
	          {"t": spec["desc"] if spec else key, "size": 13, "gap": 6},
	          {"t": "Expected: " + (spec["file"] if spec else key), "size": 9, "color": MUTED}],
	   x + 0.4, y + 0.3, w - 0.8, h - 0.4, valign="middle", align="center")


def _shadow(sp) -> None:
	sppr = sp._element.spPr
	eff = etree.SubElement(sppr, qn("a:effectLst"))
	sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="76200", dist="25400", dir="5400000",
	                      algn="bl", rotWithShape="0")
	clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
	etree.SubElement(clr, qn("a:alpha"), val="15000")


# ---------------------------------------------------------------- chevrons
def chevrons(d: Deck, s, labels: list[str], x, y, w, h, size=12, fill=NAVY, highlight=()):
	n, gap = len(labels), 0.06
	cw = (w - gap * (n - 1)) / n
	for i, lab in enumerate(labels):
		shape(s, "pentagon" if i == 0 else "chevron", x + i * (cw + gap), y, cw, h,
		      ACCENT if i in highlight else fill, line=WHITE)
		tx(d, s, lab, x + i * (cw + gap) + (0.08 if i == 0 else h * 0.5 + 0.02), y,
		   cw - h * 0.5 - 0.1 if i == 0 else cw - h - 0.04, h,
		   base={"size": size, "bold": True, "color": WHITE}, align="center", valign="middle")


# ---------------------------------------------------------------- charts
def bar_chart(d: Deck, s, x, y, w, h, categories: list[str], series: list[tuple[str, list]], *,
              horizontal=True, stacked=False, colours: list[str], label_pos="outEnd", label_size=12,
              label_bold=False, label_colour=INK, label_format=None, cat_size=11, reverse=False,
              legend=False, legend_size=11, gap=60, vmax=None, vmin=None):
	"""A bar or column chart styled as the original: no gridlines, a hidden
	value axis, value labels on the bars."""
	cd = CategoryChartData()
	cd.categories = categories
	for name, vals in series:
		cd.add_series(name, vals)
	if horizontal:
		kind = XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
	else:
		kind = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
	gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
	ch = gf.chart
	ch.has_title = False
	ch.font.name = FONT
	ch.font.size = Pt(cat_size)
	ch.font.color.rgb = rgb(INK)
	plot = ch.plots[0]
	plot.gap_width = gap
	if stacked:
		plot.overlap = 100
	plot.has_data_labels = True
	dl = plot.data_labels
	dl.font.size = Pt(label_size)
	dl.font.bold = label_bold
	dl.font.name = FONT
	dl.font.color.rgb = rgb(label_colour)
	dl.position = XL_LABEL_POSITION.CENTER if label_pos == "ctr" else XL_LABEL_POSITION.OUTSIDE_END
	dl.show_value = True
	if label_format:
		dl.number_format = label_format
		dl.number_format_is_linked = False
	for i, ser in enumerate(plot.series):
		ser.format.fill.solid()
		ser.format.fill.fore_color.rgb = rgb(colours[i % len(colours)])
		ser.invert_if_negative = False
	va = ch.value_axis
	va.visible = False
	va.has_major_gridlines = False
	va.has_minor_gridlines = False
	if vmax is not None:
		va.maximum_scale = vmax
	if vmin is not None:
		va.minimum_scale = vmin
	ca = ch.category_axis
	ca.has_major_gridlines = False
	ca.tick_labels.font.size = Pt(cat_size)
	ca.tick_labels.font.name = FONT
	ca.tick_labels.font.color.rgb = rgb(INK)
	ca.format.line.color.rgb = rgb("BFBFBF")
	if reverse:
		ca.reverse_order = True
	ch.has_legend = legend
	if legend:
		ch.legend.position = XL_LEGEND_POSITION.BOTTOM
		ch.legend.include_in_layout = False
		ch.legend.font.size = Pt(legend_size)
		ch.legend.font.name = FONT
		ch.legend.font.color.rgb = rgb(INK)
	return gf
